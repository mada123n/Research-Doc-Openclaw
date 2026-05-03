#!/usr/bin/env python3
"""generate_research_doc.py

Generate a clean, minimal Word document formatted for research reports
and summaries, per the OpenClaw research-doc specification.
"""

import argparse
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.stderr.write(
        "Error: python-docx is not installed. Install it with: pip install python-docx\n"
    )
    sys.exit(1)


DEFAULT_OUTPUT_DIR = Path.home() / ".openclaw" / "workspace" / "Research Reports"


def parse_level(value):
    try:
        level = int(value)
    except (TypeError, ValueError):
        level = 1
    return max(1, min(level, 2))


def slugify(title: str) -> str:
    slug = re.sub(r"[^\w\s-]", "", title).strip()
    slug = re.sub(r"[\s-]+", "_", slug)
    return slug or "Research_Report"


def set_run_font(run, name="Calibri", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def add_horizontal_rule(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFBFBF")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    set_run_font(run, size=10, color=(120, 120, 120))


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = r'TOC \o "1-2" \h \z \u'
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    fldChar3 = OxmlElement("w:t")
    fldChar3.text = "Right-click and choose 'Update Field' to populate the table of contents."
    fldChar4 = OxmlElement("w:fldChar")
    fldChar4.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)
    run._r.append(fldChar4)


def configure_base_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(20)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(6)

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(15)
    h2.font.bold = True
    h2.font.italic = False
    h2.font.color.rgb = RGBColor(0x3F, 0x3F, 0x3F)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)


def build_document(data, include_toc, title_arg):
    doc = Document()
    configure_base_styles(doc)

    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

    title = title_arg or data.get("title") or "Untitled Research Report"

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_p.add_run(title)
    set_run_font(title_run, size=26, bold=True)
    title_p.paragraph_format.space_after = Pt(2)

    subtitle = data.get("subtitle")
    if subtitle:
        sp = doc.add_paragraph()
        sr = sp.add_run(subtitle)
        set_run_font(sr, size=14, italic=True, color=(110, 110, 110))
        sp.paragraph_format.space_after = Pt(2)

    meta_bits = []
    author = data.get("author")
    if author:
        meta_bits.append(str(author))
    date_value = data.get("date", "auto")
    if date_value == "auto" or not date_value:
        date_value = datetime.now().strftime("%B %d, %Y")
    meta_bits.append(str(date_value))
    if meta_bits:
        mp = doc.add_paragraph()
        mr = mp.add_run("  ·  ".join(meta_bits))
        set_run_font(mr, size=10, color=(130, 130, 130))
        mp.paragraph_format.space_after = Pt(10)

    abstract = data.get("abstract")
    if abstract:
        abstract = str(abstract).strip()
    if abstract:
        ap = doc.add_paragraph()
        ar = ap.add_run(abstract)
        set_run_font(ar, size=11, italic=True, color=(70, 70, 70))
        ap.paragraph_format.space_after = Pt(4)
        rule_p = doc.add_paragraph()
        add_horizontal_rule(rule_p)
        rule_p.paragraph_format.space_after = Pt(8)

    if include_toc:
        toc_heading = doc.add_paragraph()
        thr = toc_heading.add_run("Contents")
        set_run_font(thr, size=14, bold=True)
        add_toc(doc)
        doc.add_paragraph()

    for section in data.get("sections", []) or []:
        heading = section.get("heading", "")
        level = parse_level(section.get("level", 1))
        if heading:
            doc.add_heading(str(heading), level=level)
        content = section.get("content", "")
        if content is not None and content != "":
            for para in str(content).split("\n\n"):
                para = para.strip()
                if not para:
                    continue
                p = doc.add_paragraph()
                r = p.add_run(para)
                set_run_font(r, size=12)

    references = data.get("references") or []
    if references:
        doc.add_heading("References", level=1)
        for idx, ref in enumerate(references, start=1):
            p = doc.add_paragraph()
            r = p.add_run(f"{idx}. {ref}")
            set_run_font(r, size=11)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.6)

    footer = doc.sections[0].footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number_field(fp)

    return doc, title, date_value


def print_preview(data, include_toc, title_arg):
    title = title_arg or data.get("title") or "Untitled Research Report"
    print(f"TITLE: {title}")
    if data.get("subtitle"):
        print(f"SUBTITLE: {data['subtitle']}")
    if data.get("author"):
        print(f"AUTHOR: {data['author']}")
    date_value = data.get("date", "auto")
    if date_value == "auto" or not date_value:
        date_value = datetime.now().strftime("%B %d, %Y")
    print(f"DATE: {date_value}")
    abstract = data.get("abstract")
    if abstract:
        abstract = str(abstract)
        print(f"ABSTRACT: {abstract[:120]}{'…' if len(abstract) > 120 else ''}")
    if include_toc:
        print("[Table of Contents]")
    print("---")
    for section in data.get("sections", []) or []:
        level = parse_level(section.get("level", 1))
        indent = "  " * (level - 1)
        print(f"{indent}- H{level}: {section.get('heading', '')}")
        content = str(section.get("content") or "").strip().replace("\n", " ")
        if content:
            preview = content[:80] + ("…" if len(content) > 80 else "")
            print(f"{indent}    {preview}")
    refs = data.get("references") or []
    if refs:
        print("---")
        print(f"REFERENCES ({len(refs)}):")
        for idx, ref in enumerate(refs, start=1):
            print(f"  {idx}. {ref}")


def parse_args():
    p = argparse.ArgumentParser(
        description="Generate a research-report Word document for OpenClaw."
    )
    p.add_argument("--title", help="Report title")
    p.add_argument("--content", help="Path to a JSON file containing the document content")
    p.add_argument(
        "--output-dir",
        default=str(DEFAULT_OUTPUT_DIR),
        help=f"Output directory (default: {DEFAULT_OUTPUT_DIR})",
    )
    p.add_argument("--toc", action="store_true", help="Include a table of contents")
    p.add_argument(
        "--preview",
        action="store_true",
        help="Print a plain-text outline of the document without writing a file",
    )
    return p.parse_args()


def load_content(path):
    if not path:
        return {}
    p = Path(path).expanduser()
    if not p.is_file():
        raise FileNotFoundError(f"Content file not found: {p}")
    try:
        with p.open("r", encoding="utf-8") as f:
            data = json.load(f)
    except json.JSONDecodeError as e:
        raise ValueError(f"Invalid JSON in {
